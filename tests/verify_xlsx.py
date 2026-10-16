#!/usr/bin/env python3
"""Independent check of the generated .xlsx files with openpyxl (+ LibreOffice recalculation when available).
Run `node --test tests/` first (it writes tests/out/*.xlsx and the fixture), then: python3 tests/verify_xlsx.py
Expected values are recomputed here from fixture.json by a separate implementation, not taken from the JS code.
"""
import json, os, subprocess, shutil, sys, tempfile, zipfile, re, datetime
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, 'out')
fx = json.load(open(os.path.join(OUT, 'fixture.json'))); cases = json.load(open(os.path.join(OUT, 'cases.json')))
vname = {v['id']: v['name'] for v in fx['vehicles']}
fails = []
def check(cond, msg):
    if not cond: fails.append(msg); print('  FAIL:', msg)

def select(f):
    cat = (f.get('category') or '').lower()
    ok = lambda r, key: (not f.get('vehicleId') or r['vehicleId'] == f['vehicleId']) and (not f.get('from') or r['date'] >= f['from']) and (not f.get('to') or r['date'] <= f['to']) and (not cat or r[key].lower() == cat)
    return [s for s in fx['services'] if ok(s, 'type')], [e for e in fx['expenses'] if ok(e, 'category')]

def total(s): return round(s['total'] if s['total'] is not None else (s['parts'] or 0) + (s['labor'] or 0), 2)

def summary_rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]

def recalc_with_libreoffice(path):
    """Strip cached values from formula cells, let LibreOffice recalculate, return openpyxl workbook (data_only)."""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice: return None
    tmp = tempfile.mkdtemp(); stripped = os.path.join(tmp, 'nocache.xlsx')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(stripped, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet'):
                txt = data.decode('utf8')
                txt = re.sub(r'(<c [^>]*?)( t="[a-z]+")?(>)(<f>.*?</f>)<v>.*?</v>', lambda m: m.group(1) + m.group(3) + m.group(4), txt, flags=re.S)
                data = txt.encode('utf8')
            zout.writestr(item, data)
    outdir = os.path.join(tmp, 'lo')
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir, stripped], check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(os.path.join(outdir, 'nocache.xlsx'), data_only=True)

for name, f in cases.items():
    print('==', name, f)
    path = os.path.join(OUT, name + '.xlsx')
    wb = openpyxl.load_workbook(path)                  # formulas
    wv = openpyxl.load_workbook(path, data_only=True)  # cached values
    check(wb.sheetnames[:4] == ['Vehicles', 'Maintenance Log', 'Expenses', 'Summary'], 'sheet names/order: %s' % wb.sheetnames)
    sv, ex = select(f)
    ml = list(wv['Maintenance Log'].iter_rows(min_row=2, values_only=True))
    check(len(ml) == len(sv), 'maintenance rows %d != %d' % (len(ml), len(sv)))
    check(round(sum(r[7] for r in ml), 2) == round(sum(total(s) for s in sv), 2), 'maintenance total column sum')
    check(all(isinstance(r[0], datetime.datetime) for r in ml), 'maintenance dates are real dates')
    exr = list(wv['Expenses'].iter_rows(min_row=2, values_only=True))
    check(len(exr) == len(ex), 'expense rows %d != %d' % (len(exr), len(ex)))
    check(round(sum(r[4] for r in exr), 2) == round(sum(e['amount'] for e in ex), 2), 'expense amount sum')
    vrows = list(wv['Vehicles'].iter_rows(min_row=2, values_only=True))
    check(len(vrows) == (1 if f.get('vehicleId') else 3), 'vehicle rows %d' % len(vrows))

    # Summary: find per-vehicle rows
    rows = summary_rows(wv['Summary'])
    exp_maint = round(sum(total(s) for s in sv), 2); exp_exp = round(sum(e['amount'] for e in ex), 2)
    tot_row = next(r for r in rows if r[0] == 'TOTAL')
    check(abs(tot_row[3] - exp_maint) < 0.005 and abs(tot_row[4] - exp_exp) < 0.005 and abs(tot_row[5] - (exp_maint + exp_exp)) < 0.005,
          'summary TOTAL row %s vs %s/%s' % (tot_row[:6], exp_maint, exp_exp))
    for vid, vn in vname.items():
        row = next((r for r in rows if r[0] == vn), None)
        if f.get('vehicleId') and vid != f['vehicleId']: check(row is None, vn + ' should not be in filtered summary'); continue
        m = round(sum(total(s) for s in sv if s['vehicleId'] == vid), 2); e = round(sum(x['amount'] for x in ex if x['vehicleId'] == vid), 2)
        check(row is not None and abs(row[3] - m) < 0.005 and abs(row[4] - e) < 0.005, 'per-vehicle %s: %s vs %s/%s' % (vn, row and row[:6], m, e))
    # monthly total
    months = {}
    for s in sv: months[s['date'][:7]] = months.get(s['date'][:7], 0) + total(s)
    for e in ex: months[e['date'][:7]] = months.get(e['date'][:7], 0) + e['amount']
    for r in rows:
        if isinstance(r[0], datetime.datetime) and r[3] is not None:
            k = r[0].strftime('%Y-%m'); check(abs(r[3] - months.get(k, 0)) < 0.005, 'month %s %s vs %s' % (k, r[3], months.get(k)))
    check(sum(1 for r in rows if isinstance(r[0], datetime.datetime)) == len(months), 'month rows count')
    # formulas present in the non-cached view
    fcount = sum(1 for row in wb['Summary'].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
    check(fcount >= 15, 'formula count %d' % fcount)

    # Recalculate in LibreOffice (independent formula engine) and compare to cached values
    lo = recalc_with_libreoffice(path)
    if lo is None: print('  (LibreOffice not available: recalculation check skipped)')
    else:
        a = wv['Summary']; b = lo['Summary']; bad = 0; n = 0
        for row in a.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)):
                    n += 1; v2 = b[c.coordinate].value
                    if not isinstance(v2, (int, float)) or abs(v2 - c.value) > 0.0006: bad += 1; print('   mismatch', c.coordinate, c.value, v2)
        check(bad == 0 and n > 10, 'LibreOffice recalculation matches cached values (%d numeric cells, %d mismatches)' % (n, bad))
        print('  LibreOffice recalculated %d numeric Summary cells, %d mismatches' % (n, bad))
    print('  rows: maintenance %d, expenses %d, vehicles %d' % (len(ml), len(exr), len(vrows)))
print('\nRESULT:', 'FAILED (%d)' % len(fails) if fails else 'ALL CHECKS PASSED')
sys.exit(1 if fails else 0)
